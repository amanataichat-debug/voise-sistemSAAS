"""
Agent Media Service — вложения в переписке агента обзвона с клиентами.

Каналы: личный Telegram (telegram_user_poller), Instagram DM (instagram_poller)
и WhatsApp (whatsapp_inbound, webhook Evolution API).
Что делает:
  • голосовые / аудио / видео → распознавание речи Whisper (`AGENT_STT_MODEL`,
    русский и кыргызский). Whisper не знает кыргызский как язык, поэтому язык
    не фиксируется, подсказка — `AGENT_STT_PROMPT`; если Whisper определил язык
    не из `AGENT_STT_LANGUAGES` (кыргызскую речь он часто принимает за
    казахскую/турецкую), а `ELEVENLABS_API_KEY` задан — повтор через ElevenLabs
    Scribe (batch), где кыргызский есть;
  • картинки и PDF-сканы → OCR + краткое описание vision-моделью OpenRouter
    (`AGENT_VISION_MODEL`);
  • документы (PDF с текстом, DOCX, XLSX, TXT/CSV/JSON/MD) → текст локально
    (pypdf / python-docx / openpyxl), бесплатно;
  • create_document → агент сам собирает PDF/DOCX/XLSX/TXT/CSV.

Все файлы живут в таблице agent_files (models/agent_file.py), байты — в R2
(`agent-files/...`). Распознавание и OCR списывают кредиты агента
(CreditService.charge_amount): STT — `AGENT_STT_CREDITS_PER_MINUTE` за начатую
минуту, OCR — по фактической стоимости запроса OpenRouter с наценкой
оркестратора (как обычный вызов модели).

Результат обработки превращается в строку для промпта оркестратора
(describe_for_prompt) — модель оркестратора остаётся любой, мультимодальность
ей не нужна.
"""

import asyncio
import base64
import io
import math
import mimetypes
import os
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import httpx

from backend.core.config import settings
from backend.core.logging import get_logger

logger = get_logger(__name__)

OPENAI_TRANSCRIBE_URL = "https://api.openai.com/v1/audio/transcriptions"
ELEVEN_STT_URL = "https://api.elevenlabs.io/v1/speech-to-text"
ELEVEN_STT_MODEL = os.getenv("AGENT_STT_SCRIBE_MODEL", "scribe_v2")

# Whisper принимает файлы до 25 МБ.
WHISPER_MAX_BYTES = 25 * 1024 * 1024
# Сколько текста храним и сколько показываем оркестратору в самом сообщении.
MAX_STORED_TEXT = 200_000
PROMPT_PREVIEW_CHARS = 1500
# PDF: если текстового слоя меньше — это скан, отправляем в OCR.
PDF_MIN_TEXT_CHARS = 80
PDF_OCR_MAX_BYTES = 15 * 1024 * 1024

FONTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "assets", "fonts")

TEXT_EXTS = {".txt", ".csv", ".json", ".md", ".html", ".htm", ".xml", ".log"}
DOC_FORMATS = ("pdf", "docx", "xlsx", "txt", "csv")

KIND_LABELS = {
    "voice": "🎤 Голосовое сообщение",
    "audio": "🎵 Аудио",
    "video": "🎥 Видео",
    "image": "🖼 Изображение",
    "document": "📄 Документ",
    "other": "📎 Файл",
}


# ============================================================================
# КЛАССИФИКАЦИЯ
# ============================================================================

def guess_mime(filename: str, mime: Optional[str] = None) -> str:
    if mime and mime != "application/octet-stream":
        return mime
    guessed, _ = mimetypes.guess_type(filename or "")
    return guessed or (mime or "application/octet-stream")


def ext_for(mime: str, filename: str = "") -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext:
        return ext
    special = {
        "audio/ogg": ".ogg", "audio/opus": ".ogg", "audio/mpeg": ".mp3",
        "audio/mp4": ".m4a", "audio/x-m4a": ".m4a", "audio/wav": ".wav",
        "audio/x-wav": ".wav", "audio/webm": ".webm", "video/mp4": ".mp4",
        "image/jpeg": ".jpg", "image/png": ".png", "image/webp": ".webp",
        "application/pdf": ".pdf",
    }
    return special.get((mime or "").split(";")[0].strip(), mimetypes.guess_extension(mime or "") or "")


def classify(mime: str, filename: str = "", hint: Optional[str] = None) -> str:
    """Вид вложения: voice / audio / video / image / document / other."""
    if hint in ("voice", "audio", "video", "image", "document"):
        return hint
    m = (mime or "").lower()
    ext = os.path.splitext(filename or "")[1].lower()
    if m.startswith("audio/"):
        return "audio"
    if m.startswith("video/"):
        return "video"
    if m.startswith("image/"):
        return "image"
    if m in ("application/pdf",) or ext in (".pdf", ".docx", ".xlsx", ".xlsm") or ext in TEXT_EXTS \
            or m.startswith("text/") or "wordprocessingml" in m or "spreadsheetml" in m:
        return "document"
    return "other"


def fmt_duration(seconds: Optional[float]) -> str:
    if not seconds:
        return ""
    s = int(round(seconds))
    return f"{s // 60}:{s % 60:02d}"


# ============================================================================
# РАСПОЗНАВАНИЕ РЕЧИ (Whisper, ru + ky; Scribe как запасной для кыргызского)
# ============================================================================

def _allowed_languages() -> set:
    return {x.strip().lower() for x in (settings.AGENT_STT_LANGUAGES or "").split(",") if x.strip()}


async def _whisper(data: bytes, filename: str, mime: str) -> Dict[str, Any]:
    if not settings.OPENAI_API_KEY:
        raise RuntimeError("OPENAI_API_KEY not configured")
    form = {
        "model": settings.AGENT_STT_MODEL,
        "response_format": "verbose_json",
    }
    if settings.AGENT_STT_PROMPT:
        form["prompt"] = settings.AGENT_STT_PROMPT
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            OPENAI_TRANSCRIBE_URL,
            headers={"Authorization": f"Bearer {settings.OPENAI_API_KEY}"},
            data=form,
            files={"file": (filename, data, mime)},
        )
    if resp.status_code != 200:
        raise RuntimeError(f"whisper {resp.status_code}: {resp.text[:300]}")
    out = resp.json()
    return {
        "text": (out.get("text") or "").strip(),
        "language": (out.get("language") or "").lower(),
        "duration": float(out.get("duration") or 0) or None,
        "provider": settings.AGENT_STT_MODEL,
    }


async def _scribe(data: bytes, filename: str, mime: str) -> Dict[str, Any]:
    if not settings.ELEVENLABS_API_KEY:
        raise RuntimeError("ELEVENLABS_API_KEY not configured")
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            ELEVEN_STT_URL,
            headers={"xi-api-key": settings.ELEVENLABS_API_KEY},
            data={"model_id": ELEVEN_STT_MODEL, "tag_audio_events": "false"},
            files={"file": (filename, data, mime)},
        )
    if resp.status_code != 200:
        raise RuntimeError(f"scribe {resp.status_code}: {resp.text[:300]}")
    out = resp.json()
    duration = None
    words = out.get("words") or []
    if words and isinstance(words[-1], dict):
        try:
            duration = float(words[-1].get("end") or 0) or None
        except (TypeError, ValueError):
            duration = None
    return {
        "text": (out.get("text") or "").strip(),
        "language": (out.get("language_code") or "").lower(),
        "duration": duration,
        "provider": ELEVEN_STT_MODEL,
    }


async def transcribe(data: bytes, filename: str, mime: str) -> Dict[str, Any]:
    """
    Распознать речь. Возвращает {ok, text, language, duration, provider} или
    {ok: False, error}. Whisper — основной, Scribe — когда Whisper упал или
    определил язык не из AGENT_STT_LANGUAGES.
    """
    allowed = _allowed_languages()
    scribe_ok = settings.AGENT_STT_SCRIBE_FALLBACK and bool(settings.ELEVENLABS_API_KEY)
    whisper_res = None
    try:
        whisper_res = await _whisper(data, filename, mime)
    except Exception as e:
        logger.warning(f"[AGENT-MEDIA] whisper failed: {e}")

    need_scribe = whisper_res is None or (
        whisper_res["text"] and whisper_res["language"] and whisper_res["language"] not in allowed
    )
    if need_scribe and scribe_ok:
        try:
            res = await _scribe(data, filename, mime)
            if whisper_res and not res.get("duration"):
                res["duration"] = whisper_res.get("duration")
            logger.info(
                f"[AGENT-MEDIA] scribe used (whisper lang={whisper_res['language'] if whisper_res else 'n/a'}) "
                f"→ lang={res['language']}"
            )
            return {"ok": True, **res}
        except Exception as e:
            logger.warning(f"[AGENT-MEDIA] scribe failed: {e}")

    if whisper_res is not None:
        return {"ok": True, **whisper_res}
    return {"ok": False, "error": "transcription_failed"}


# ============================================================================
# OCR / VISION (OpenRouter)
# ============================================================================

VISION_PROMPT = (
    "Клиент прислал это вложение в мессенджере менеджеру по продажам. "
    "1) Перепиши дословно весь читаемый текст (сохрани числа, суммы, даты, "
    "реквизиты, номера телефонов). 2) Одним-двумя предложениями опиши, что это "
    "(фото товара, скриншот оплаты, чек, паспорт, договор и т.п.). Пиши на языке "
    "документа, описание — по-русски. Формат ответа строго:\n"
    "ТЕКСТ: <текст или «нет»>\nОПИСАНИЕ: <описание>"
)


async def vision_extract(data: bytes, mime: str, filename: str) -> Dict[str, Any]:
    """
    OCR + описание картинки или PDF-скана. Возвращает
    {ok, text, model, prompt_tokens, completion_tokens, cost_usd} или {ok: False, error}.
    """
    from backend.services.openrouter_client import get_openrouter_client

    if not settings.OPENROUTER_API_KEY:
        return {"ok": False, "error": "openrouter_not_configured"}
    b64 = base64.b64encode(data).decode()
    extra = None
    if mime == "application/pdf":
        part = {
            "type": "file",
            "file": {"filename": filename or "document.pdf", "file_data": f"data:application/pdf;base64,{b64}"},
        }
        extra = {"plugins": [{"id": "file-parser", "pdf": {"engine": "native"}}]}
    else:
        part = {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}}

    model = settings.AGENT_VISION_MODEL
    try:
        resp = await get_openrouter_client().chat_completion(
            model=model,
            messages=[{"role": "user", "content": [{"type": "text", "text": VISION_PROMPT}, part]}],
            tools=None,
            temperature=0.0,
            max_tokens=4000,
            timeout=120.0,
            extra_payload=extra,
        )
    except Exception as e:
        logger.warning(f"[AGENT-MEDIA] vision failed ({model}): {e}")
        return {"ok": False, "error": f"vision_failed: {e}"}
    usage = resp.get("usage") or {}
    text = ((resp.get("choices") or [{}])[0].get("message") or {}).get("content") or ""
    if isinstance(text, list):
        text = "".join(p.get("text", "") for p in text if isinstance(p, dict))
    return {
        "ok": True,
        "text": text.strip(),
        "model": model,
        "prompt_tokens": int(usage.get("prompt_tokens", 0) or 0),
        "completion_tokens": int(usage.get("completion_tokens", 0) or 0),
        "cost_usd": usage.get("cost"),
    }


# ============================================================================
# ИЗВЛЕЧЕНИЕ ТЕКСТА ИЗ ДОКУМЕНТОВ (локально, без кредитов)
# ============================================================================

def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_document_text(data: bytes, filename: str, mime: str) -> Tuple[Optional[str], Optional[int]]:
    """
    Текст документа (синхронно — звать через asyncio.to_thread).
    Возвращает (text, pages) — text None, если формат не поддерживается.
    """
    ext = os.path.splitext(filename or "")[1].lower()
    m = (mime or "").lower()
    if ext == ".pdf" or m == "application/pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                parts.append("")
        return "\n".join(p.strip() for p in parts if p.strip()), len(reader.pages)
    if ext == ".docx" or "wordprocessingml" in m:
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return "\n".join(lines), None
    if ext in (".xlsx", ".xlsm") or "spreadsheetml" in m:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets[:5]:
            lines.append(f"# Лист: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 500:
                    lines.append("…")
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    lines.append(" | ".join(cells).rstrip(" |"))
        return "\n".join(lines), None
    if ext in TEXT_EXTS or m.startswith("text/") or m == "application/json":
        return _decode_text(data), None
    return None, None


# ============================================================================
# БИЛЛИНГ
# ============================================================================

def stt_credits(duration_seconds: Optional[float]) -> int:
    minutes = max(1, math.ceil((duration_seconds or 1) / 60.0))
    return minutes * max(0, settings.AGENT_STT_CREDITS_PER_MINUTE)


def vision_credits(res: Dict[str, Any]) -> int:
    from backend.services.agent_models import _CREDITS_PER_USD
    from backend.services.credit_service import CreditService
    cost = res.get("cost_usd")
    try:
        if cost is not None and float(cost) > 0:
            return max(1, math.ceil(float(cost) * _CREDITS_PER_USD))
    except (TypeError, ValueError):
        pass
    return CreditService.calculate_cost(res.get("model") or "", res.get("prompt_tokens", 0), res.get("completion_tokens", 0))


def _charge(db, row, amount: int, ref_type: str, model_slug: Optional[str], notes: str) -> None:
    from backend.services.credit_service import CreditService
    if amount <= 0:
        return
    try:
        CreditService.charge_amount(
            db, row.user_id, amount, ref_type=ref_type, ref_id=row.id,
            model_slug=model_slug, notes=notes,
        )
        row.credits_charged = int(row.credits_charged or 0) + amount
    except Exception as e:
        logger.error(f"[AGENT-MEDIA] charge failed for file {row.id}: {e}")
        try:
            db.rollback()
        except Exception:
            pass


def _has_credits(db, user_id) -> bool:
    from backend.models.user import User
    user = db.query(User).filter(User.id == user_id).first()
    return bool(user) and (user.credits_balance or 0) > 0


# ============================================================================
# ХРАНЕНИЕ
# ============================================================================

def _storage_key(agent_config_id, source: str, file_id, ext: str) -> str:
    now = datetime.utcnow()
    return f"agent-files/{agent_config_id}/{source}/{now.year}/{now.month:02d}/{file_id}{ext}"


async def store_file(
    db,
    *,
    user_id,
    agent_config_id,
    data: bytes,
    filename: str,
    mime: Optional[str] = None,
    source: str = "library",
    kind: Optional[str] = None,
    agent_contact_id=None,
    channel: Optional[str] = None,
    external_message_id: Optional[str] = None,
    title: Optional[str] = None,
    description: Optional[str] = None,
    status: str = "processing",
):
    """Создать AgentFile и положить байты в R2 (без R2 файл сохраняется только метаданными)."""
    from backend.models.agent_file import AgentFile
    from backend.services.r2_storage import R2StorageService

    mime = guess_mime(filename, mime)
    filename = (filename or "file").strip()[:255] or "file"
    file_id = uuid.uuid4()
    row = AgentFile(
        id=file_id,
        user_id=user_id,
        agent_config_id=agent_config_id,
        agent_contact_id=agent_contact_id,
        source=source,
        channel=channel,
        external_message_id=(str(external_message_id) if external_message_id else None),
        kind=kind or classify(mime, filename),
        filename=filename,
        mime_type=mime[:128],
        size_bytes=len(data or b""),
        title=(title or None),
        description=(description or None),
        status=status,
    )
    if data and R2StorageService.objects_available():
        key = _storage_key(agent_config_id, source, file_id, ext_for(mime, filename))
        if await R2StorageService.put_object(key, data, mime):
            row.storage_key = key
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


async def load_bytes(row) -> Optional[bytes]:
    from backend.services.r2_storage import R2StorageService
    if not row or not row.storage_key:
        return None
    return await R2StorageService.get_object(row.storage_key)


async def delete_file(db, row) -> None:
    from backend.services.r2_storage import R2StorageService
    if row.storage_key:
        await R2StorageService.delete_object(row.storage_key)
    db.delete(row)
    db.commit()


def file_link(row, expires_seconds: int = 7 * 24 * 3600) -> Optional[str]:
    from backend.services.r2_storage import R2StorageService
    if not row or not row.storage_key:
        return None
    return R2StorageService.presigned_url(row.storage_key, expires_seconds, filename=row.filename)


# ============================================================================
# ОБРАБОТКА ВЛОЖЕНИЯ
# ============================================================================

async def analyze_file(db, row, data: bytes, charge: bool = True) -> None:
    """
    Распознать/прочитать файл и записать результат в row.extracted_text.
    Кредиты списываются только за платные шаги (STT, OCR). Ошибки — в row.error,
    наружу не бросает.
    """
    kind = row.kind
    mime = row.mime_type or "application/octet-stream"
    filename = row.filename or "file"
    max_bytes = settings.AGENT_MEDIA_MAX_MB * 1024 * 1024

    try:
        if not data:
            row.status, row.error = "failed", "empty_file"
            return
        if len(data) > max_bytes:
            row.status, row.error = "skipped", "file_too_large"
            return

        if kind in ("voice", "audio", "video"):
            if len(data) > WHISPER_MAX_BYTES:
                row.status, row.error = "skipped", "file_too_large"
                return
            if row.duration_seconds and row.duration_seconds > settings.AGENT_VOICE_MAX_SECONDS:
                row.status, row.error = "skipped", "too_long"
                return
            if charge and not _has_credits(db, row.user_id):
                row.status, row.error = "skipped", "insufficient_credits"
                return
            if not filename.lower().endswith((".ogg", ".mp3", ".mp4", ".m4a", ".wav", ".webm", ".mpeg", ".mpga", ".flac", ".oga")):
                filename = f"{filename}{ext_for(mime) or '.ogg'}"
            res = await transcribe(data, filename, mime)
            if not res.get("ok"):
                row.status, row.error = "failed", res.get("error")
                return
            row.extracted_text = res["text"][:MAX_STORED_TEXT]
            if res.get("duration"):
                row.duration_seconds = res["duration"]
            row.status = "ready"
            if charge:
                _charge(db, row, stt_credits(row.duration_seconds), "agent_media_stt",
                        res.get("provider"), f"{kind} {fmt_duration(row.duration_seconds)} lang={res.get('language') or '?'}")
            return

        if kind == "document":
            text, pages = await asyncio.to_thread(extract_document_text, data, filename, mime)
            is_pdf = mime == "application/pdf" or filename.lower().endswith(".pdf")
            if is_pdf and len((text or "").strip()) < PDF_MIN_TEXT_CHARS and len(data) <= PDF_OCR_MAX_BYTES:
                # Скан без текстового слоя → OCR
                if charge and not _has_credits(db, row.user_id):
                    row.status, row.error = "skipped", "insufficient_credits"
                    return
                res = await vision_extract(data, "application/pdf", filename)
                if res.get("ok"):
                    row.extracted_text = res["text"][:MAX_STORED_TEXT]
                    row.status = "ready"
                    if charge:
                        _charge(db, row, vision_credits(res), "agent_media_ocr", res.get("model"), f"pdf ocr {filename}")
                    return
                row.status, row.error = "failed", res.get("error")
                return
            if text is None:
                row.status, row.error = "skipped", "unsupported_format"
                return
            row.extracted_text = text[:MAX_STORED_TEXT]
            row.status = "ready" if text.strip() else "failed"
            if not text.strip():
                row.error = "no_text"
            return

        if kind == "image":
            if charge and not _has_credits(db, row.user_id):
                row.status, row.error = "skipped", "insufficient_credits"
                return
            res = await vision_extract(data, mime, filename)
            if not res.get("ok"):
                row.status, row.error = "failed", res.get("error")
                return
            row.extracted_text = res["text"][:MAX_STORED_TEXT]
            row.status = "ready"
            if charge:
                _charge(db, row, vision_credits(res), "agent_media_ocr", res.get("model"), f"image {filename}")
            return

        row.status, row.error = "skipped", "unsupported_format"
    except Exception as e:
        logger.error(f"[AGENT-MEDIA] analyze_file {row.id} failed: {e}", exc_info=True)
        row.status, row.error = "failed", f"{type(e).__name__}: {e}"[:500]
    finally:
        try:
            db.commit()
        except Exception:
            db.rollback()


async def process_inbound(
    db,
    *,
    user_id,
    agent_config_id,
    agent_contact_id,
    channel: str,
    external_message_id,
    data: Optional[bytes],
    filename: str,
    mime: Optional[str],
    kind_hint: Optional[str] = None,
    duration_hint: Optional[float] = None,
    error: Optional[str] = None,
):
    """
    Вложение клиента из канала → AgentFile (inbound) с распознанным текстом.
    error — причина, по которой поллер не смог скачать файл (сохраняется как
    failed/skipped без обработки).
    """
    mime = guess_mime(filename, mime)
    kind = classify(mime, filename, kind_hint)
    row = await store_file(
        db,
        user_id=user_id,
        agent_config_id=agent_config_id,
        data=data or b"",
        filename=filename or f"{kind}{ext_for(mime)}",
        mime=mime,
        source="inbound",
        kind=kind,
        agent_contact_id=agent_contact_id,
        channel=channel,
        external_message_id=external_message_id,
    )
    if duration_hint:
        row.duration_seconds = float(duration_hint)
    if error:
        row.status = "skipped" if error == "file_too_large" else "failed"
        row.error = error
        db.commit()
        return row
    await analyze_file(db, row, data or b"")
    logger.info(
        f"[AGENT-MEDIA] {channel} {kind} {row.filename} → {row.status}"
        f"{' (' + row.error + ')' if row.error else ''}, {row.credits_charged} cr"
    )
    return row


ERROR_HUMAN = {
    "file_too_large": "файл слишком большой",
    "too_long": "слишком длинное",
    "insufficient_credits": "не обработано — закончились кредиты",
    "unsupported_format": "формат не поддерживается",
    "transcription_failed": "не удалось распознать речь",
    "no_text": "текста не найдено",
    "download_failed": "не удалось скачать файл",
    "empty_file": "пустой файл",
}


def describe_for_prompt(row, caption: Optional[str] = None, preview_chars: int = PROMPT_PREVIEW_CHARS) -> str:
    """
    Строка вложения для переписки и промпта оркестратора. Содержит file_id,
    чтобы агент мог перечитать полный текст (read_agent_file) или переслать файл.
    """
    label = KIND_LABELS.get(row.kind, KIND_LABELS["other"])
    meta = []
    if row.kind in ("voice", "audio", "video") and row.duration_seconds:
        meta.append(fmt_duration(row.duration_seconds))
    if row.kind in ("document", "other", "image") and row.filename and not row.filename.startswith(("photo", "image")):
        meta.append(f"«{row.filename}»")
    head = label + (f" ({', '.join(meta)})" if meta else "")
    head += f" [file_id={row.id}]"

    text = (row.extracted_text or "").strip()
    if row.status == "ready" and text:
        short = text if len(text) <= preview_chars else text[:preview_chars].rstrip() + "… (полностью — read_agent_file)"
        if row.kind in ("voice", "audio", "video"):
            body = f": «{short}»"
        else:
            body = f":\n{short}"
    else:
        reason = ERROR_HUMAN.get((row.error or "").split(":")[0], row.error or "без текста")
        body = f" — {reason}"
    out = head + body
    if caption:
        out += f"\nПодпись клиента: {caption.strip()}"
    return out


# ============================================================================
# СОЗДАНИЕ ДОКУМЕНТОВ (create_document)
# ============================================================================

def _safe_filename(title: str, ext: str) -> str:
    base = "".join(ch if (ch.isalnum() or ch in " -_.()") else "_" for ch in (title or "document")).strip()
    base = " ".join(base.split())[:80] or "document"
    return f"{base}.{ext}"


def _lines(content: str) -> List[str]:
    return (content or "").replace("\r\n", "\n").split("\n")


def _render_pdf(title: str, content: str, table: Optional[List[List[Any]]]) -> bytes:
    import logging
    from fpdf import FPDF

    # fpdf2 сабсетит шрифт через fontTools, который пишет в INFO десятки строк.
    logging.getLogger("fontTools").setLevel(logging.WARNING)

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("DejaVu", "", os.path.join(FONTS_DIR, "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", os.path.join(FONTS_DIR, "DejaVuSans-Bold.ttf"))
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    if title:
        pdf.set_font("DejaVu", "B", 16)
        pdf.multi_cell(width, 9, title)
        pdf.ln(3)
    for line in _lines(content):
        stripped = line.strip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            pdf.set_font("DejaVu", "B", 14 if level <= 1 else 12)
            pdf.multi_cell(width, 7, stripped.lstrip("#").strip())
            pdf.ln(1)
            continue
        pdf.set_font("DejaVu", "", 11)
        if stripped.startswith(("- ", "* ", "• ")):
            pdf.multi_cell(width, 6, "• " + stripped[2:].strip())
        elif stripped:
            pdf.multi_cell(width, 6, stripped.replace("**", ""))
        else:
            pdf.ln(4)
    if table:
        pdf.ln(4)
        pdf.set_font("DejaVu", "", 10)
        rows = [["" if c is None else str(c) for c in r] for r in table if isinstance(r, (list, tuple))]
        with pdf.table(first_row_as_headings=True) as t:
            ncols = max(len(r) for r in rows) if rows else 0
            for r in rows:
                rr = t.row()
                for c in r + [""] * (ncols - len(r)):
                    rr.cell(c)
    return bytes(pdf.output())


def _render_docx(title: str, content: str, table: Optional[List[List[Any]]]) -> bytes:
    import docx
    d = docx.Document()
    if title:
        d.add_heading(title, level=0)
    for line in _lines(content):
        stripped = line.strip()
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            d.add_heading(stripped.lstrip("#").strip(), level=min(max(level, 1), 3))
        elif stripped.startswith(("- ", "* ", "• ")):
            d.add_paragraph(stripped[2:].strip(), style="List Bullet")
        elif stripped:
            d.add_paragraph(stripped.replace("**", ""))
    if table:
        rows = [["" if c is None else str(c) for c in r] for r in table if isinstance(r, (list, tuple))]
        if rows:
            ncols = max(len(r) for r in rows)
            t = d.add_table(rows=len(rows), cols=ncols)
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _render_xlsx(title: str, content: str, table: Optional[List[List[Any]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Лист1")[:31].replace("/", "-")
    rows = table or [[c.strip() for c in ln.split(";")] for ln in _lines(content) if ln.strip()]
    for r in rows:
        ws.append(list(r) if isinstance(r, (list, tuple)) else [str(r)])
    if rows:
        for cell in ws[1]:
            cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _render_csv(title: str, content: str, table: Optional[List[List[Any]]]) -> bytes:
    import csv
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    if table:
        for r in table:
            w.writerow(["" if c is None else c for c in (r if isinstance(r, (list, tuple)) else [r])])
    else:
        for ln in _lines(content):
            if ln.strip():
                w.writerow([c.strip() for c in ln.split(";")])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _render_txt(title: str, content: str, table: Optional[List[List[Any]]]) -> bytes:
    parts = [title, ""] if title else []
    parts.append(content or "")
    if table:
        parts.append("")
        parts.extend(" | ".join("" if c is None else str(c) for c in r) for r in table if isinstance(r, (list, tuple)))
    return "\n".join(parts).encode("utf-8")


RENDERERS = {
    "pdf": (_render_pdf, "application/pdf"),
    "docx": (_render_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "xlsx": (_render_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "csv": (_render_csv, "text/csv"),
    "txt": (_render_txt, "text/plain"),
}


async def create_document(
    db,
    *,
    user_id,
    agent_config_id,
    title: str,
    content: str,
    fmt: str = "pdf",
    table: Optional[List[List[Any]]] = None,
    agent_contact_id=None,
    description: Optional[str] = None,
):
    """Собрать документ и сохранить как AgentFile(source=generated). Бросает ValueError."""
    fmt = (fmt or "pdf").lower().lstrip(".")
    if fmt not in RENDERERS:
        raise ValueError(f"Формат {fmt} не поддерживается: {', '.join(DOC_FORMATS)}")
    if not (content or "").strip() and not table:
        raise ValueError("Пустой документ: передай content или table")
    renderer, mime = RENDERERS[fmt]
    data = await asyncio.to_thread(renderer, (title or "").strip(), content or "", table)
    filename = _safe_filename(title or "document", fmt)
    row = await store_file(
        db,
        user_id=user_id,
        agent_config_id=agent_config_id,
        data=data,
        filename=filename,
        mime=mime,
        source="generated",
        kind="document",
        agent_contact_id=agent_contact_id,
        title=(title or filename)[:255],
        description=description,
        status="ready",
    )
    text = (content or "").strip()
    if table:
        text += "\n" + "\n".join(" | ".join("" if c is None else str(c) for c in r) for r in table if isinstance(r, (list, tuple)))
    row.extracted_text = text[:MAX_STORED_TEXT]
    db.commit()
    return row
